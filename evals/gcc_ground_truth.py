#!/usr/bin/env python3
"""Validate engine output against the GCC workbooks' ground-truth recon_state.

Each GCC dataset (gcc-data/) ships a Recon_Sales sheet whose ``recon_state`` /
``sub_bucket_code`` is the intended reconciliation outcome per document. This
harness runs the engine on the converted CSVs (data/gcc/<ISO3>/, produced by
gcc-data/to_engine_csv.py) and checks that the engine reproduces the outcomes an
amount-based three-way engine is *designed* to determine:

  * clean matches (OK)                       -> a matched unit (cell E1/E4/E5)
  * booked but never e-invoiced (ME-NONE,
    GEN-FAIL, ING-FAIL)                       -> E2  BILLED_UNREPORTED
  * e-invoiced but never booked (MG-UNBK)     -> E7  REPORTED_UNBOOKED

These three are asserted at 100%. The remaining ground-truth sub-buckets are
business-rule overlays on top of the amounts - late reporting (ME-LATE), a live
e-invoice with a later booking (MG-LIVE), buyer-TIN / metadata / clearance-status
diffs (MM-TIN, MM-META, MM-STAT) and out-of-scope B2C/exempt supplies (OOS) - and
carry equal or matchable amounts. An engine that matches on [tax, net, gross]
legitimately does not surface them from the numbers alone, so they are reported
in a confusion matrix rather than scored. OOS in particular lands as E2 here:
the engine has no OOS rule, so an out-of-scope supply with no e-invoice reads as
a genuine reporting gap - a real finding worth surfacing.

Exit status is non-zero if any asserted target is missed or conservation fails.
"""
from __future__ import annotations

import os
import sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import openpyxl  # noqa: E402

from core import config as config_mod, pipeline  # noqa: E402

# (country pack, ISO3 dataset id)
DATASETS = [("OM", "OMN"), ("SA", "KSA"), ("AE", "UAE")]
CLIENT = "gcc-einvoice"

# Sub-buckets the amount-based engine must reproduce, and the coarse engine
# outcome each one requires.
ASSERTED = {
    "OK": "MATCHED",
    "ME-NONE": "BILLED_UNREPORTED",
    "GEN-FAIL": "BILLED_UNREPORTED",
    "ING-FAIL": "BILLED_UNREPORTED",
    "MG-UNBK": "REPORTED_UNBOOKED",
}

RECON_STATE_COARSE = {
    "exact_match": "MATCHED",
    "mismatch": "MISMATCH",
    "missing_in_einvoice": "BILLED_UNREPORTED",
    "missing_in_gl": "REPORTED_UNBOOKED",
    "out_of_scope": "OUT_OF_SCOPE",
}


def _rows(ws):
    hdr = [h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if h is not None]
    for r in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(hdr, r))


def _ground_truth(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    gt = {r["document_no"]: {"state": r["recon_state"], "sub": r["sub_bucket_code"]}
          for r in _rows(wb["Recon_Sales"])}
    wb.close()
    return gt


def _engine_class(cells: set) -> str:
    """Coarse existence outcome for a document from its unit cells."""
    if cells & {"E1", "E4", "E5"}:
        return "MATCHED"
    if cells == {"E2"}:
        return "BILLED_UNREPORTED"
    if cells == {"E7"}:
        return "REPORTED_UNBOOKED"
    if cells == {"E2", "E7"}:
        return "SPLIT_UNMATCHED"
    if cells <= {"E6"}:
        return "RECEIVABLE_ONLY"
    return "/".join(sorted(cells)) or "NONE"


def evaluate(country, iso3):
    cfg = config_mod.load_all(
        os.path.join(ROOT, "config", "engine.default.json"),
        os.path.join(ROOT, "config", "country-packs", f"{country}.pack.json"),
        os.path.join(ROOT, "config", "client-profiles", f"{CLIENT}.profile.json"),
    )
    ddir = os.path.join(ROOT, "data", "gcc", iso3)
    res = pipeline.run(cfg, os.path.join(ddir, "gl.csv"), os.path.join(ddir, "ar.csv"),
                       os.path.join(ddir, "einv.csv"), period_filter=None)

    unit_by_member = {rid: u for u in res["units"] for rid in u["member_ids"]}
    doc_cells = defaultdict(set)
    for d in res["docs"]:
        num = d.get("billing_doc_id") or d.get("accounting_doc_id")
        u = unit_by_member.get(d["row_id"])
        if u:
            doc_cells[num].add(u["existence_cell"])

    gt = _ground_truth(os.path.join(ROOT, "gcc-data", f"{iso3}_einvoice_recon_dataset.xlsx"))
    records = []
    for doc, g in gt.items():
        got = _engine_class(doc_cells.get(doc, set()))
        records.append({"doc": doc, "state": g["state"], "sub": g["sub"], "got": got})
    return res, records


def main():
    all_records = []
    conservation_ok = True
    for country, iso3 in DATASETS:
        res, records = evaluate(country, iso3)
        all_records += records
        conservation_ok = conservation_ok and res["conservation"]["ok"]
        matched = sum(1 for r in records if r["got"] == "MATCHED")
        print(f"{iso3} (pack {country}): {len(records)} ground-truth docs, "
              f"conservation={'PASS' if res['conservation']['ok'] else 'FAIL'}, "
              f"matched={matched}")

    # ---- asserted sub-buckets -----------------------------------------------
    print(f"\n{'='*70}\nASSERTED SUB-BUCKETS (amount/existence-decidable)\n{'='*70}")
    all_pass = True
    for sub, expected in ASSERTED.items():
        recs = [r for r in all_records if r["sub"] == sub]
        ok = sum(1 for r in recs if r["got"] == expected)
        rate = ok / len(recs) if recs else 1.0
        passed = rate >= 1.0
        all_pass = all_pass and passed
        print(f"  {'PASS' if passed else 'FAIL'}  {sub:10} -> {expected:18} "
              f"{ok}/{len(recs)}  ({rate:.0%})")

    # ---- transparency confusion matrix --------------------------------------
    print(f"\n{'='*70}\nCONFUSION: recon_state -> engine outcome (count)\n{'='*70}")
    conf = Counter((RECON_STATE_COARSE.get(r["state"], r["state"]), r["got"]) for r in all_records)
    for (exp, got), n in sorted(conf.items(), key=lambda x: -x[1]):
        mark = "" if exp == got else "   (overlay / not amount-decidable)"
        print(f"  {exp:18} -> {got:18} : {n}{mark}")

    print(f"\n{'='*70}\nSUB-BUCKET BREAKDOWN (informational)\n{'='*70}")
    sub_conf = defaultdict(Counter)
    for r in all_records:
        sub_conf[r["sub"]][r["got"]] += 1
    for sub in sorted(sub_conf):
        tag = "" if sub in ASSERTED else "  (overlay)"
        print(f"  {sub:10} {dict(sub_conf[sub])}{tag}")

    print(f"\n  {'PASS' if conservation_ok else 'FAIL'}  conservation_check (all datasets)")
    all_pass = all_pass and conservation_ok
    print(f"\n{'='*70}\nBUILD {'PASS' if all_pass else 'FAIL'}\n{'='*70}")
    return 0 if all_pass else 1


if __name__ == "__main__":
    raise SystemExit(main())
