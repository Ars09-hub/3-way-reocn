#!/usr/bin/env python3
"""Reshape a GCC e-invoice recon workbook into the engine's CSV contract.

The GCC datasets (see gcc-data/run_index.json) are doc-level: one summarised
row per invoice carrying taxable_value and tax_value together, with no separate
AR extract and no GL line detail. The reconciliation engine consumes three
line-/doc-level CSVs instead:

  * gl.csv   - line-level GL postings (customer-control / revenue / tax lines)
  * ar.csv   - one customer (AR) document per invoice
  * einv.csv - one cleared e-invoice per row

This converter derives all three from the workbook's *sales* side
(GL_Sales + Einvoice_Sales), which maps onto the engine's R/C/E output-VAT
model. The purchase side is a distinct input-tax reconciliation and is out of
scope for this engine.

Column names emitted here match config/client-profiles/gcc-einvoice.profile.json.

Usage:
  python gcc-data/to_engine_csv.py                # convert all datasets in run_index.json
  python gcc-data/to_engine_csv.py --iso3 OMN     # a single country
"""
from __future__ import annotations

import argparse
import csv
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# GL account numbers we synthesise the three posting lines against. These fall
# inside the account_classes ranges of the gcc-einvoice client profile.
ACCT_CUSTOMER_CONTROL = "140000"
ACCT_REVENUE = "410000"
ACCT_TAX_OUTPUT = "230000"

# Standard-rate tax code per jurisdiction (the country pack maps the code -> rate).
STANDARD_CODE = {"UAE": "S5", "KSA": "S15", "OMN": "S5"}
TAX_CODE_MAP = {"S": None, "Z": "Z0", "E": "EX"}  # S resolved per-country below

GL_COLUMNS = [
    "entity_id", "accounting_doc_id", "fiscal_year", "line_no", "gl_account",
    "doc_type", "doc_date", "posting_date", "currency", "fx_rate",
    "amount_signed", "tax_code", "customer_id", "customer_tin", "customer_name",
]
AR_COLUMNS = [
    "entity_id", "billing_doc_id", "accounting_doc_id", "official_doc_id",
    "fiscal_year", "doc_type", "doc_date", "posting_date", "currency", "fx_rate",
    "net_amount", "tax_amount", "gross_amount", "tax_code",
    "customer_id", "customer_tin", "customer_name",
]
EINV_COLUMNS = [
    "entity_id", "invoice_number", "seller_tin", "document_type", "issue_date",
    "clearance_date", "status", "uuid", "currency", "taxable_amount",
    "tax_amount", "total_amount", "tax_code", "buyer_tin", "buyer_name",
    "is_consolidated",
]


def _sheet_rows(ws):
    header = [h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if h is not None]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(header, row))


def _tax_code(iso3: str, raw) -> str:
    raw = (str(raw).strip() if raw is not None else "") or "S"
    if raw == "S":
        return STANDARD_CODE.get(iso3, "S5")
    return TAX_CODE_MAP.get(raw, raw)


def _year(dstr) -> str:
    return str(dstr)[:4] if dstr else ""


def convert(iso3: str, xlsx_path: str, out_dir: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    gl_rows = list(_sheet_rows(wb["GL_Sales"]))
    einv_rows = list(_sheet_rows(wb["Einvoice_Sales"]))
    wb.close()

    # seller_tin -> entity_id, learned from the GL side so e-invoices (which carry
    # no company code) can be blocked/matched by entity like the GL and AR docs.
    tin_to_entity = {}
    for r in gl_rows:
        if r.get("seller_tin"):
            tin_to_entity.setdefault(str(r["seller_tin"]), r.get("entity_id"))

    gl_out, ar_out = [], []
    for r in gl_rows:
        doc = r.get("document_no")
        entity = r.get("entity_id")
        code = _tax_code(iso3, r.get("tax_code"))
        taxable = r.get("taxable_value") or 0
        tax = r.get("tax_value") or 0
        gross = round(float(taxable) + float(tax), 3)
        fy = _year(r.get("posting_date") or r.get("voucher_date"))
        common = {
            "entity_id": entity, "doc_type": r.get("doc_type"),
            "doc_date": r.get("voucher_date"), "posting_date": r.get("posting_date"),
            "currency": r.get("currency"), "fx_rate": "1.0", "tax_code": code,
            "customer_id": r.get("customer_tin"), "customer_tin": r.get("customer_tin"),
            "customer_name": r.get("customer_name"),
        }
        # GL: customer-control (debit gross) / revenue (credit net) / tax (credit tax).
        lines = [(ACCT_CUSTOMER_CONTROL, gross), (ACCT_REVENUE, -float(taxable))]
        if float(tax) != 0:
            lines.append((ACCT_TAX_OUTPUT, -float(tax)))
        for i, (acct, amt) in enumerate(lines, start=1):
            gl_out.append({**common, "accounting_doc_id": doc, "fiscal_year": fy,
                           "line_no": i, "gl_account": acct,
                           "amount_signed": f"{amt:.3f}"})
        # AR: one customer document, AWKEY-linked to the GL voucher.
        ar_out.append({**common, "billing_doc_id": doc, "accounting_doc_id": doc,
                       "official_doc_id": doc, "fiscal_year": fy,
                       "net_amount": f"{float(taxable):.3f}", "tax_amount": f"{float(tax):.3f}",
                       "gross_amount": f"{gross:.3f}"})

    einv_out = []
    for r in einv_rows:
        seller = r.get("seller_trn") or r.get("seller_vat_number")
        buyer = r.get("buyer_trn") or r.get("buyer_vat_number")
        itc = r.get("invoice_type_code")
        doc_type = "Credit Note" if str(itc) == "381" else "Invoice"
        # status / timestamp / uuid column names vary by jurisdiction.
        clr = (r.get("clearance_timestamp") or r.get("status_timestamp")
               or r.get("reporting_timestamp"))
        uuid = r.get("uuid") or r.get("peppol_message_id")
        einv_out.append({
            "entity_id": tin_to_entity.get(str(seller), ""),
            "invoice_number": r.get("invoice_number"),
            "seller_tin": seller, "document_type": doc_type,
            "issue_date": r.get("issue_date"), "clearance_date": clr,
            "status": "CLEARED", "uuid": uuid,
            "currency": r.get("document_currency_code") or r.get("currency"),
            "taxable_amount": r.get("tax_exclusive_amount") or r.get("line_extension_amount"),
            "tax_amount": r.get("tax_amount"),
            "total_amount": r.get("tax_inclusive_amount") or r.get("payable_amount"),
            "tax_code": _tax_code(iso3, r.get("tax_category_code")),
            "buyer_tin": buyer, "buyer_name": r.get("buyer_name"),
            "is_consolidated": "false",
        })

    os.makedirs(out_dir, exist_ok=True)
    _write(os.path.join(out_dir, "gl.csv"), GL_COLUMNS, gl_out)
    _write(os.path.join(out_dir, "ar.csv"), AR_COLUMNS, ar_out)
    _write(os.path.join(out_dir, "einv.csv"), EINV_COLUMNS, einv_out)
    return {"gl_lines": len(gl_out), "ar_docs": len(ar_out), "einv_docs": len(einv_out),
            "out_dir": out_dir}


def _write(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=columns)
        w.writeheader()
        for r in rows:
            w.writerow({c: ("" if r.get(c) is None else r.get(c)) for c in columns})


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--iso3", help="convert only this dataset (e.g. OMN); default all")
    ap.add_argument("--index", default=os.path.join(HERE, "run_index.json"))
    ap.add_argument("--out-root", default=os.path.join(ROOT, "data", "gcc"))
    args = ap.parse_args(argv)

    index = json.load(open(args.index))
    for ds in index["datasets"]:
        iso3 = ds["country_iso3"]
        if args.iso3 and iso3 != args.iso3:
            continue
        xlsx = os.path.join(ROOT, ds["file"])
        out_dir = os.path.join(args.out_root, iso3)
        stats = convert(iso3, xlsx, out_dir)
        print(f"{iso3}: gl_lines={stats['gl_lines']} ar_docs={stats['ar_docs']} "
              f"einv_docs={stats['einv_docs']} -> {os.path.relpath(out_dir, ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
